# -*- coding: utf-8 -*-
"""
KANS 越南直播间 · 主播排名 —— 数据内核
读线上 Google 表 -> 合并多个直播间 tab -> 计算当日/当月排名 -> 生成三语 Excel。

关键：把当月的 SKINCARE 直播间 和 OFFICIAL + KHTH 直播间 两个 tab 合并，
      同一个主播的 GMV 跨直播间加总。每月自动找当月的 tab，无需手动改。
被界面程序 main_app.py 调用；也可单独 python3 report_core.py 跑。
"""

import csv
import io
import os
import re
import ssl
import sys
import urllib.request
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================
# ★★★ 配置区 ★★★
# 默认值对应 KANS 越南的表；换成自己的表时，在本目录放一个 config.json 覆盖任意字段：
#   {"workbook_id": "...", "room_keywords": ["..."], "vnd_per_rmb": 3860,
#    "output_dir": "~/Desktop/主播排名报表", "columns": {"date":0,"host_vn":7,"host_cn":8,"gmv":13,"roi":16}}
# 也可用环境变量 KANS_WORKBOOK_ID 单独指定表 ID。
# ============================================================
import json as _json

_DEFAULTS = {
    "workbook_id": "",   # 线上 Google 表（发布为 CSV）
    "room_keywords": ["SKINCARE", "OFFICIAL", "KHTH"],   # tab 名含其一即纳入合并（不区分大小写）
    "vnd_per_rmb": 3860,                                  # 1 元人民币 = 多少越南盾
    "output_dir": "~/Desktop/KANS主播排名报表",           # 报表输出目录
    "columns": {"date": 0, "host_vn": 7, "host_cn": 8, "gmv": 13, "roi": 16},  # 列索引，从 0 开始
    # 本地货币：code 用在英文/本地语工作表表头，zh 用在中文表备注
    "currency": {"code": "VND", "zh": "越南盾"},
    # 第三张工作表的本地语言（默认越南语）。换国家把这几段文案换成当地语言即可
    "local_lang": {
        "sheet": "Tiếng Việt",
        "headers": ["Hạng", "Host", "GMV trước hoàn ({cur})"],
        "day_title": "Xếp hạng ngày  {d:%d/%m/%Y}",
        "month_title": "Xếp hạng tháng  {d:%m/%Y}",
        "note": "Phạm vi: dữ liệu đến hết {d:%d/%m/%Y} (không tính hôm nay) | GMV trước hoàn | tỷ giá 1 CNY = {rate:,} {cur}",
    },
}

def _load_cfg():
    cfg = dict(_DEFAULTS)
    here = os.path.dirname(os.path.abspath(__file__))
    for cand in (os.path.join(here, "config.json"), os.path.join(os.path.dirname(getattr(sys, "executable", "")), "config.json")):
        if cand and os.path.isfile(cand):
            try:
                user = _json.load(open(cand, encoding="utf-8"))
                cols = {**cfg["columns"], **user.pop("columns", {})}
                cur = {**cfg["currency"], **user.pop("currency", {})}
                loc = {**cfg["local_lang"], **user.pop("local_lang", {})}
                cfg.update(user); cfg["columns"] = cols; cfg["currency"] = cur; cfg["local_lang"] = loc
                break
            except Exception:  # noqa
                pass
    if os.environ.get("KANS_WORKBOOK_ID"):
        cfg["workbook_id"] = os.environ["KANS_WORKBOOK_ID"]
    return cfg

_CFG = _load_cfg()
WORKBOOK_ID = _CFG["workbook_id"]
ROOM_KEYWORDS = _CFG["room_keywords"]
VND_PER_RMB = _CFG["vnd_per_rmb"]
OUTPUT_DIR = os.path.expanduser(_CFG["output_dir"])
COL_DATE = _CFG["columns"]["date"]
COL_HOST_VN = _CFG["columns"]["host_vn"]
COL_HOST_CN = _CFG["columns"]["host_cn"]
COL_GMV_HOST = _CFG["columns"]["gmv"]     # 退前 GMV（越南盾）
COL_ROI = _CFG["columns"]["roi"]
CUR = _CFG["currency"]["code"]          # 本地货币代码，如 VND / THB / IDR
CUR_ZH = _CFG["currency"]["zh"]        # 本地货币中文名
_LOC = _CFG["local_lang"]
# ============================================================


# ---------------------- 网络 ----------------------
def _ssl_contexts():
    """系统默认 -> certifi -> 免校验（公开只读链接兜底），保证换机也能抓。"""
    yield ssl.create_default_context()
    try:
        import certifi
        yield ssl.create_default_context(cafile=certifi.where())
    except Exception:  # noqa
        pass
    yield ssl._create_unverified_context()


def _fetch_text(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    last = None
    for ctx in _ssl_contexts():
        try:
            with urllib.request.urlopen(req, timeout=60, context=ctx) as resp:
                return resp.read().decode("utf-8", errors="replace")
        except (urllib.error.URLError, ssl.SSLError) as e:
            last = e
    raise last


def _list_all_tabs():
    """从 htmlview 拿到工作簿里所有 tab 的 (名字, gid)。"""
    url = f"https://docs.google.com/spreadsheets/d/{WORKBOOK_ID}/htmlview"
    html = _fetch_text(url)
    pairs = re.findall(r'name:\s*"((?:[^"\\]|\\.)*)"[^{}]*?gid:\s*"(\d+)"', html)
    out = []
    for name, gid in pairs:
        if "\\u" in name:
            try:
                name = name.encode().decode("unicode_escape")
            except Exception:  # noqa
                pass
        out.append((name, gid))
    return out


def discover_room_tabs(month):
    """找出当月要合并的直播间 tab：tab 名含 'Tháng {月}' 且含直播间关键词。"""
    month_re = re.compile(rf"th[aá]ng\s*0*{month}\b", re.IGNORECASE)
    tabs = []
    for name, gid in _list_all_tabs():
        if not month_re.search(name):
            continue
        if any(kw.lower() in name.lower() for kw in ROOM_KEYWORDS):
            tabs.append((name.strip(), gid))
    return tabs


def _fetch_csv_rows(gid):
    url = (f"https://docs.google.com/spreadsheets/d/{WORKBOOK_ID}"
           f"/export?format=csv&gid={gid}")
    return list(csv.reader(io.StringIO(_fetch_text(url))))


# ---------------------- 解析 ----------------------
def parse_vnd(text):
    if text is None:
        return None
    t = text.strip().replace(".", "").replace(" ", "").replace(" ", "")
    if t in ("", "-", "#N/A"):
        return None
    if not t.lstrip("-").isdigit():
        return None
    return int(t)


def parse_date(text):
    try:
        return datetime.strptime(text.strip(), "%d/%m/%Y")
    except (ValueError, AttributeError):
        return None


def vn_display_name(host_vn):
    parts = host_vn.strip().split()
    if parts and parts[-1].isdigit():
        parts = parts[:-1]
    return " ".join(parts).strip()


def _parse_rows(rows):
    out = []
    for row in rows[1:]:
        if len(row) <= COL_ROI:
            continue
        d = parse_date(row[COL_DATE])
        if d is None:
            continue
        vn = vn_display_name(row[COL_HOST_VN])
        if not vn:               # 未来排班的占位空行
            continue
        gmv = parse_vnd(row[COL_GMV_HOST])
        if gmv is None:
            continue
        cn = row[COL_HOST_CN].strip()
        if cn in ("", "#N/A"):
            cn = vn              # 中文名没匹配上就用越文名兜底
        out.append({"date": d, "vn": vn, "cn": cn, "gmv": gmv})
    return out


def load_sessions(reference_date, progress=None):
    """按参照日期所在【月】，找到当月各直播间 tab，全部抓下来合并。"""
    def say(m):
        if progress:
            progress(m)
    tabs = discover_room_tabs(reference_date.month)
    if not tabs:
        raise RuntimeError("NO_TABS")
    say("TABS:" + " + ".join(n for n, _ in tabs))
    sessions = []
    for name, gid in tabs:
        sessions.extend(_parse_rows(_fetch_csv_rows(gid)))
    return sessions, tabs


def rank(sessions):
    """按主播汇总 GMV（同名主播跨直播间自动加总），降序。"""
    agg = {}
    for s in sessions:
        a = agg.setdefault(s["vn"],
                           {"vn": s["vn"], "cn": s["cn"], "gmv": 0, "sessions": 0})
        a["gmv"] += s["gmv"]
        a["sessions"] += 1
    return sorted(agg.values(), key=lambda x: x["gmv"], reverse=True)


# --------------------------- Excel ---------------------------
HEADER_FILL = PatternFill("solid", fgColor="BDD7EE")
TITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="9AA7B4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")

SHEET_LANGS = {
    "简体中文": {
        "name_key": "cn",
        "to_money": lambda v: round(v / VND_PER_RMB),
        "headers": ["排名", "主播", "退前GMV(元)"],
        "day_title": lambda d: f"当日排名  {d:%Y-%m-%d}",
        "month_title": lambda d: f"当月总排名  {d:%Y-%m}",
        "note": lambda d: (f"统计口径：数据截止 {d:%Y-%m-%d}（不含当天）｜ "
                           f"退前GMV ｜ 汇率 1元 = {VND_PER_RMB:,} {CUR_ZH}"),
        "widths": [8, 18, 18],
    },
    "English": {
        "name_key": "vn",
        "to_money": lambda v: v,
        "headers": ["Rank", "Host", f"GMV before refund ({CUR})"],
        "day_title": lambda d: f"Daily Ranking  {d:%Y-%m-%d}",
        "month_title": lambda d: f"Monthly Ranking  {d:%Y-%m}",
        "note": lambda d: (f"Scope: data through {d:%Y-%m-%d} (today excluded) | "
                           f"GMV before refund | rate 1 CNY = {VND_PER_RMB:,} {CUR}"),
        "widths": [8, 20, 26],
    },
    _LOC["sheet"]: {
        "name_key": "vn",
        "to_money": lambda v: v,
        "headers": [h.replace("{cur}", CUR) for h in _LOC["headers"]],
        "day_title": lambda d: _LOC["day_title"].format(d=d),
        "month_title": lambda d: _LOC["month_title"].format(d=d),
        "note": lambda d: _LOC["note"].format(d=d, rate=VND_PER_RMB, cur=CUR),
        "widths": [8, 20, 24],
    },
}


def _write_block(ws, start_row, title, cfg, records):
    ncol = len(cfg["headers"])
    ws.merge_cells(start_row=start_row, start_column=1,
                   end_row=start_row, end_column=ncol)
    c = ws.cell(start_row, 1, title)
    c.font = Font(bold=True, size=13)
    c.alignment = CENTER
    c.fill = TITLE_FILL
    for col in range(1, ncol + 1):
        ws.cell(start_row, col).border = BORDER

    hr = start_row + 1
    for col, h in enumerate(cfg["headers"], 1):
        cell = ws.cell(hr, col, h)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER

    for i, rec in enumerate(records, 1):
        r = hr + i
        vals = [i, rec[cfg["name_key"]], cfg["to_money"](rec["gmv"])]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(r, col, v)
            cell.border = BORDER
            cell.alignment = LEFT if col == 2 else CENTER
            if col == 3:
                cell.number_format = "#,##0"
    return hr + len(records) + 2


def _build_sheet(ws, cfg, day_rank, month_rank, day_date):
    for i, w in enumerate(cfg["widths"], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.sheet_view.showGridLines = False
    # 顶部口径说明
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=len(cfg["headers"]))
    note = ws.cell(1, 1, cfg["note"](day_date))
    note.font = Font(size=9, italic=True, color="9A3B3F")
    note.alignment = LEFT
    nxt = _write_block(ws, 3, cfg["day_title"](day_date), cfg, day_rank)
    _write_block(ws, nxt, cfg["month_title"](day_date), cfg, month_rank)


def generate_report(target_date_str=None, progress=None):
    """
    target_date_str: None=自动出前一天；"2026-07-05"=补出指定日期。
    progress: 可选回调 progress(msg)。
    返回 dict：out_path / day_date / day_count / month_count / n_day / n_month / rooms。
    """
    def say(m):
        if progress:
            progress(m)

    # 先定“参照日期”，用来决定读哪个月的 tab
    if target_date_str:
        reference = datetime.strptime(target_date_str.strip(), "%Y-%m-%d")
    else:
        reference = datetime.now() - timedelta(days=1)   # 前一天所在月

    say("FETCHING")
    sessions, tabs = load_sessions(reference, progress=progress)
    if not sessions:
        raise RuntimeError("EMPTY")

    if target_date_str:
        target = reference
        day = [s for s in sessions if s["date"].date() == target.date()]
        month = [s for s in sessions
                 if s["date"].year == target.year
                 and s["date"].month == target.month
                 and s["date"].date() <= target.date()]
    else:
        today = datetime.now().date()          # 按运行电脑真实日期
        past = [s["date"] for s in sessions
                if s["date"].date() < today and s["gmv"] > 0]
        if not past:
            raise RuntimeError("NO_PAST_DATA")
        target = max(past)                      # 前一天（当天不统计）
        day = [s for s in sessions if s["date"].date() == target.date()]
        month = [s for s in sessions
                 if s["date"].year == target.year
                 and s["date"].month == target.month
                 and s["date"].date() < today]

    day_rank, month_rank = rank(day), rank(month)

    wb = Workbook()
    first = True
    for sheet_name, cfg in SHEET_LANGS.items():
        ws = wb.active if first else wb.create_sheet()
        ws.title = sheet_name
        first = False
        _build_sheet(ws, cfg, day_rank, month_rank, target)

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    fname = f"KANS主播排名_{target:%Y-%m-%d}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, fname)
    wb.save(out_path)

    return {
        "out_path": out_path,
        "day_date": target,
        "day_count": len(day),
        "month_count": len(month),
        "n_day": len(day_rank),
        "n_month": len(month_rank),
        "rooms": [n for n, _ in tabs],
    }


if __name__ == "__main__":
    r = generate_report(progress=lambda m: print("...", m))
    print("合并直播间:", " + ".join(r["rooms"]))
    print("当日", r["day_date"].strftime("%Y-%m-%d"),
          r["day_count"], "场", r["n_day"], "主播")
    print("当月", r["month_count"], "场", r["n_month"], "主播")
    print("已生成", r["out_path"])
